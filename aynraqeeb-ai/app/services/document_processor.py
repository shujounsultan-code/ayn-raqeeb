import io
import fitz  # PyMuPDF
import pdfplumber
import docx
from langchain_text_splitters import RecursiveCharacterTextSplitter
from app.config import settings


class DocumentProcessor:
    """محرك معالجة معقد (Complex Parser) — يدعم استخراج الجداول والنصوص بذكاء"""

    def __init__(self):
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP,
            separators=["\n\n", "\n", ".", "،", "؟", "!", " ", ""],
        )

    def extract_text(self, file_bytes: bytes, file_type: str) -> str:
        """نقطة الدخول الرئيسية لاستخراج النص"""
        if file_type == "pdf":
            return self._extract_complex_pdf(file_bytes)
        elif file_type == "docx":
            return self._extract_from_docx(file_bytes)
        elif file_type == "txt":
            return file_bytes.decode("utf-8", errors="ignore")
        else:
            raise ValueError(f"نوع ملف غير مدعوم: {file_type}")

    def _extract_complex_pdf(self, file_bytes: bytes) -> str:
        """
        محرك معالجة PDF متقدم (Mini DeepDoc):
        - يستخدم PyMuPDF للنصوص السريعة.
        - يستخدم pdfplumber للجداول المعقدة.
        """
        full_content = []
        
        # نفتح الملف باستخدام كلا المكتبتين
        pdf_stream = io.BytesIO(file_bytes)
        
        with pdfplumber.open(pdf_stream) as plub_pdf:
            # نستخدم PyMuPDF أيضاً للسرعة في النصوص
            doc_fitz = fitz.open(stream=file_bytes, filetype="pdf")
            
            for page_num in range(len(plub_pdf.pages)):
                page_text = ""
                
                # 1. استخراج الجداول أولاً من الصفحة
                plub_page = plub_pdf.pages[page_num]
                tables = plub_page.extract_tables()
                
                # 2. استخراج النص العادي من الصفحة
                fitz_page = doc_fitz[page_num]
                page_text = fitz_page.get_text()
                
                # 3. إذا وجدت جداول، نحولها لـ Markdown ونضيفها
                if tables:
                    for table in tables:
                        md_table = self._table_to_markdown(table)
                        if md_table:
                            page_text += f"\n\n[جدول مستخرج]:\n{md_table}\n"
                
                full_content.append(page_text)
                
            doc_fitz.close()
            
        return "\n\n--- صفحة جديدة ---\n\n".join(full_content)

    def _table_to_markdown(self, table: list[list]) -> str:
        """تحويل مصفوفة الجدول إلى صيغة Markdown ليفهمها الـ LLM"""
        if not table or not any(table):
            return ""
        
        # تنظيف البيانات من None وخطوط فارغة
        clean_table = [[(str(cell) if cell is not None else "") for cell in row] for row in table if any(row)]
        if not clean_table: return ""

        headers = clean_table[0]
        rows = clean_table[1:] if len(clean_table) > 1 else []
        
        # بناء الـ Markdown
        md = "| " + " | ".join(headers) + " |\n"
        md += "| " + " | ".join(["---"] * len(headers)) + " |\n"
        for row in rows:
            # التأكد من توافق عدد خلايا الصف مع الهيدر
            if len(row) < len(headers):
                row.extend([""] * (len(headers) - len(row)))
            md += "| " + " | ".join(row[:len(headers)]) + " |\n"
            
        return md

    def _extract_from_docx(self, file_bytes: bytes) -> str:
        """استخراج النص والجداول من ملفات Word"""
        doc = docx.Document(io.BytesIO(file_bytes))
        content = []
        
        # استخراج النصوص
        for para in doc.paragraphs:
            if para.text.strip():
                content.append(para.text)
        
        # استخراج الجداول من Word وتحويلها لـ Markdown
        for table in doc.tables:
            data = []
            for row in table.rows:
                data.append([cell.text.strip() for cell in row.cells])
            md_table = self._table_to_markdown(data)
            if md_table:
                content.append(f"\n[جدول مستخرج من Word]:\n{md_table}\n")
                
        return "\n".join(content)

    def chunk_text(self, text: str) -> list[str]:
        """تقطيع النص إلى chunks متينة"""
        return self.splitter.split_text(text)


document_processor = DocumentProcessor()
